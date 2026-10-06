"""XLSX hidden sheet adapter — stores chunks in hidden Excel worksheets.

Each chunk becomes an .xlsx with data in a VeryHidden worksheet:
    xlsx://{output_dir}/{chunk_id}.xlsx

Requires: pip install picard[xlsx]  (openpyxl)
"""

from __future__ import annotations

import asyncio
import uuid
from pathlib import Path
from typing import Sequence

from picard.adapters.base import Adapter, AdapterProfile, StoredChunkRef
from picard.chunker import Chunk


class XlsxHiddenAdapter(Adapter):
    """Hide chunks in XLSX VeryHidden worksheets."""

    def __init__(self, output_dir: str | Path, *, max_chunk_bytes: int = 65_536) -> None:
        self._dir = Path(output_dir)
        self._max = max_chunk_bytes

    def profile(self) -> AdapterProfile:
        return AdapterProfile(
            name="xlsx-hidden",
            max_chunk_bytes=self._max,
            persistent=True,
            ordered=False,
            rate_limit_per_min=120,
        )

    async def write_chunk(self, encoded_chunk: str) -> StoredChunkRef:
        try:
            chunk = Chunk.decode(encoded_chunk)
            chunk_id = f"{chunk.manifest_id}_{chunk.block_index}_{chunk.chunk_index}"
        except ValueError:
            chunk_id = uuid.uuid4().hex

        path = self._dir / f"{chunk_id}.xlsx"

        def _write():
            from openpyxl import Workbook

            self._dir.mkdir(parents=True, exist_ok=True)
            wb = Workbook()
            visible = wb.active
            visible.title = "Sheet1"
            visible["A1"] = "Report"

            hidden = wb.create_sheet("__picard")
            hidden.sheet_state = "veryHidden"

            row_size = 32000
            for i in range(0, len(encoded_chunk), row_size):
                hidden.cell(row=i // row_size + 1, column=1, value=encoded_chunk[i:i + row_size])

            wb.save(path)

        await asyncio.to_thread(_write)
        return StoredChunkRef(adapter_name="xlsx-hidden", location=f"xlsx://{path}")

    async def read_chunk(self, ref: StoredChunkRef) -> str:
        path = Path(ref.location.split("://", 1)[1])

        def _read():
            from openpyxl import load_workbook

            wb = load_workbook(path)
            hidden = wb["__picard"]
            parts = []
            for row in hidden.iter_rows(min_col=1, max_col=1, values_only=True):
                if row[0] is not None:
                    parts.append(str(row[0]))
            return "".join(parts)

        return await asyncio.to_thread(_read)

    async def delete_chunk(self, ref: StoredChunkRef) -> bool:
        path = Path(ref.location.split("://", 1)[1])
        try:
            path.unlink()
            return True
        except FileNotFoundError:
            return False

    async def list_chunks(self, manifest_id: str) -> Sequence[StoredChunkRef]:
        refs: list[StoredChunkRef] = []
        if self._dir.is_dir():
            for p in self._dir.glob(f"{manifest_id}_*.xlsx"):
                refs.append(
                    StoredChunkRef(adapter_name="xlsx-hidden", location=f"xlsx://{p}")
                )
        return refs

    async def health_check(self) -> bool:
        return self._dir.is_dir() or not self._dir.exists()

    async def capacity(self) -> dict[str, int | None]:
        import shutil

        if not self._dir.exists():
            return {"used_bytes": 0, "free_bytes": None, "total_bytes": None, "chunk_count": 0}
        usage = shutil.disk_usage(self._dir)
        count = sum(1 for _ in self._dir.glob("*.xlsx"))
        used = sum(f.stat().st_size for f in self._dir.glob("*.xlsx"))
        return {"used_bytes": used, "free_bytes": usage.free, "total_bytes": usage.total, "chunk_count": count}
